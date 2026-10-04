"""CSV / TSV / XLSX: rows become table blocks that keep their header row."""

from __future__ import annotations

import csv
from collections.abc import Iterator
from itertools import islice
from pathlib import Path
from typing import TYPE_CHECKING

from src.core.errors import ParseError, ProviderUnavailableError
from src.parsing.language import detect_language
from src.parsing.markdown import table_block
from src.parsing.text import read_utf8
from src.ports.parsing import ParsedUnit

if TYPE_CHECKING:
    from src.core.config import Settings

ROWS_PER_UNIT = 50


class CsvParser:
    name = "csv"
    extensions = frozenset({".csv", ".tsv"})

    def __init__(self, settings: Settings | None = None) -> None:
        pass

    def iter_units(self, path: Path) -> Iterator[ParsedUnit]:
        delimiter = "\t" if path.suffix.lower() == ".tsv" else ","
        # Decode up front so a bad encoding is a ParseError rather than a mid-stream crash.
        reader = csv.reader(read_utf8(path).splitlines(), delimiter=delimiter)
        try:
            header = next(reader)
        except StopIteration:
            return
        language = None
        number = 0
        while rows := list(islice(reader, ROWS_PER_UNIT)):
            if language is None:  # detected once per file: per-block detection is far too slow on large files
                language = detect_language(" ".join(header + [c for r in rows for c in r]))
            number += 1
            yield ParsedUnit(unit=number, tables=[table_block(header, rows, language)])


class XlsxParser:
    name = "xlsx"
    extensions = frozenset({".xlsx", ".xlsm"})

    def __init__(self, settings: Settings | None = None) -> None:
        pass

    def iter_units(self, path: Path) -> Iterator[ParsedUnit]:
        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise ProviderUnavailableError(
                "XLSX parsing needs openpyxl: pip install 'turinton-rag[xlsx]'"
            ) from exc
        try:
            workbook = load_workbook(path, read_only=True, data_only=True)
        except Exception as exc:  # openpyxl raises zipfile/XML errors of many types
            raise ParseError(f"{path.name}: cannot open workbook: {exc}") from exc
        number = 0
        try:
            for sheet in workbook.worksheets:
                rows = sheet.iter_rows(values_only=True)
                header = next(rows, None)
                if header is None:
                    continue
                language = None
                while block := [r for r in islice(rows, ROWS_PER_UNIT) if any(v is not None for v in r)]:
                    if language is None:
                        language = detect_language(
                            " ".join(str(v) for r in (header, *block) for v in r if v is not None)
                        )
                    number += 1
                    yield ParsedUnit(
                        unit=number,
                        tables=[table_block(header, block, language, label=f"Sheet '{sheet.title}'")],
                    )
        finally:
            workbook.close()
