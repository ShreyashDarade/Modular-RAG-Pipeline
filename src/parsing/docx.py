from __future__ import annotations

import hashlib
from collections.abc import Iterator
from pathlib import Path
from typing import TYPE_CHECKING

from src.core.errors import ParseError, ProviderUnavailableError
from src.parsing.language import detect_language
from src.parsing.markdown import table_block
from src.ports.parsing import ImageRef, ParsedUnit, TextBlock

if TYPE_CHECKING:
    from src.core.config import Settings

_SECTION_STYLES = ("Title", "Heading 1", "Heading 2")


class DocxParser:
    """python-docx. A new unit starts at every Title / Heading 1 / Heading 2; tables keep their
    header row and images are attached to the unit that contains them."""

    name = "docx"
    extensions = frozenset({".docx"})

    def __init__(self, settings: Settings | None = None) -> None:
        self._min_side = settings.ocr_min_image_side if settings is not None else 0

    def iter_units(self, path: Path) -> Iterator[ParsedUnit]:
        try:
            import docx
            from docx.table import Table
            from docx.text.paragraph import Paragraph
        except ImportError as exc:
            raise ProviderUnavailableError(
                "DOCX parsing needs python-docx: pip install 'ai-rag-info[docx]'"
            ) from exc
        try:
            document = docx.Document(str(path))
        except Exception as exc:  # python-docx raises PackageNotFoundError, BadZipFile, KeyError, ...
            raise ParseError(f"{path.name}: cannot open document: {exc}") from exc

        number = 0
        lines: list[str] = []
        unit = ParsedUnit(unit=1)
        seen_images: set[str] = set()
        seen_content: set[str] = set()

        def flush() -> ParsedUnit | None:
            text = "\n".join(line for line in lines if line.strip())
            if text:
                unit.texts.append(TextBlock(text, detect_language(text)))
            return unit if (unit.texts or unit.tables or unit.images) else None

        for child in document.element.body.iterchildren():
            tag = child.tag.rsplit("}", 1)[-1]
            if tag == "p":
                paragraph = Paragraph(child, document)
                style = paragraph.style.name if paragraph.style is not None else ""
                if style in _SECTION_STYLES and (lines or unit.tables or unit.images):
                    finished = flush()
                    if finished is not None:
                        number += 1
                        finished.unit = number
                        yield finished
                    lines, unit = [], ParsedUnit(unit=number + 1)
                lines.append(paragraph.text)
                for rel_id in child.xpath(".//a:blip/@r:embed"):
                    if rel_id in seen_images:
                        continue
                    seen_images.add(rel_id)
                    blob = document.part.related_parts[rel_id].blob
                    digest = hashlib.sha256(blob).hexdigest()
                    if digest in seen_content:
                        continue
                    seen_content.add(digest)
                    unit.images.append(ImageRef(label=f"{path.name}#{rel_id}", data=blob))
            elif tag == "tbl":
                rows = [[cell.text for cell in row.cells] for row in Table(child, document).rows]
                if rows:
                    header, *body = rows
                    unit.tables.append(
                        table_block(header, body, detect_language(" ".join(c for r in rows for c in r)))
                    )
        finished = flush()
        if finished is not None:
            finished.unit = number + 1
            yield finished
