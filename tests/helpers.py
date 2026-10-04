"""Builders for real test documents."""

from __future__ import annotations

from collections.abc import Sequence
from pathlib import Path

import cv2
import numpy as np


def make_pdf(
    path: Path,
    pages: Sequence[str],
    *,
    table: Sequence[Sequence[str]] | None = None,
    image_on_page: int | None = None,
    image_png: bytes | None = None,
) -> Path:
    import pymupdf

    doc = pymupdf.open()
    for number, text in enumerate(pages, start=1):
        page = doc.new_page()
        y = 72
        for line in text.split("\n"):
            page.insert_text((72, y), line, fontsize=11)
            y += 16
        if table and number == 1:
            x0, y0, cw, rh = 72, y + 20, 100, 24
            rows, cols = len(table), len(table[0])
            for r in range(rows + 1):
                page.draw_line((x0, y0 + r * rh), (x0 + cols * cw, y0 + r * rh))
            for c in range(cols + 1):
                page.draw_line((x0 + c * cw, y0), (x0 + c * cw, y0 + rows * rh))
            for r, row in enumerate(table):
                for c, value in enumerate(row):
                    page.insert_text((x0 + c * cw + 6, y0 + r * rh + 16), value, fontsize=10)
        if image_png is not None and image_on_page == number:
            page.insert_image(pymupdf.Rect(72, 400, 372, 520), stream=image_png)
    doc.save(path)
    doc.close()
    return path


def make_docx(path: Path, sections: dict[str, str]) -> Path:
    import docx

    document = docx.Document()
    for heading, body in sections.items():
        document.add_heading(heading, 1)
        document.add_paragraph(body)
    document.save(path)
    return path


def text_png(lines: Sequence[str], *, angle: float = 0.0) -> bytes:
    """A noisy scan-like page of text, optionally rotated."""
    rng = np.random.default_rng(7)
    image = np.full((120 + 70 * len(lines), 1000, 3), 245, np.uint8)
    for i, line in enumerate(lines):
        cv2.putText(image, line, (50, 90 + i * 70), cv2.FONT_HERSHEY_SIMPLEX, 1.1, (30, 30, 30), 2)
    image = np.clip(image.astype(int) + rng.integers(-6, 7, image.shape), 0, 255).astype(np.uint8)
    if angle:
        h, w = image.shape[:2]
        matrix = cv2.getRotationMatrix2D((w / 2, h / 2), angle, 1.0)
        image = cv2.warpAffine(image, matrix, (w, h), borderValue=(245, 245, 245))
    return cv2.imencode(".png", image)[1].tobytes()
