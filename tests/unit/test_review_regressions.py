"""Regression tests for defects found in review. Each one failed before its fix."""

from __future__ import annotations

import asyncio
from pathlib import Path

import httpx
import openpyxl
import pytest
from fastapi import FastAPI, File, UploadFile
from pydantic import ValidationError
from src.api.middleware import MaxBodySizeMiddleware
from src.chat.service import ChatService
from src.chat.stores import MemoryConversationStore
from src.core.bootstrap import build_registries
from src.core.config import Settings
from src.core.errors import QueueFullError, RequestTimeoutError, UpstreamError
from src.core.types import ChatMessage, JobSpec
from src.jobs.inprocess import InProcessJobBackend
from src.models.adapters import CachingEmbedder, LangChainEmbedder
from src.parsing.registry import ParserSet
from src.runtime.cache import MemoryCache
from src.runtime.concurrency import Bulkhead, deadline
from src.runtime.redis_client import new_client


def parse(path: Path):
    settings = Settings(_env_file=None)
    parser_set = ParserSet(build_registries(settings), settings)
    return list(parser_set.for_path(path).iter_units(path))


# --- tabular parsers ------------------------------------------------------------------------
def test_csv_cells_may_contain_newlines(tmp_path):
    (tmp_path / "a.csv").write_text('name,address\nAda,"1 Main St\nSpringfield"\nBob,"x\r\ny"\n', newline="")
    [unit] = parse(tmp_path / "a.csv")
    assert unit.tables[0].row_count == 2, "a multi-line cell is one cell of one row, not two rows"
    assert "1 Main St Springfield" in unit.tables[0].markdown, "the line break becomes a space"
    assert "StSpringfield" not in unit.tables[0].markdown, "lines must not be glued together"


def test_xlsx_rows_after_a_long_blank_gap_are_not_dropped(tmp_path):
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.append(["id", "value"])
    for i in range(50):
        sheet.append([i, i])
    for _ in range(120):  # spacer rows: whole 50-row slices of nothing
        sheet.append([None, None])
    for i in range(1000, 1010):
        sheet.append([i, i])
    workbook.save(tmp_path / "gap.xlsx")
    assert sum(u.tables[0].row_count for u in parse(tmp_path / "gap.xlsx")) == 60


# --- in-process queue -----------------------------------------------------------------------
async def test_a_retry_is_never_refused_when_the_queue_is_full():
    """Regression: re-queueing a failed job into a full queue raised QueueFull inside the worker
    loop, which killed every worker and left the job 'running' forever."""
    backend = InProcessJobBackend(Settings(_env_file=None, ingest_queue_max_size=1, ingest_max_attempts=3))
    spec = JobSpec(collection="c", path="/x")
    first = await backend.submit(spec)  # the queue is now at capacity
    with pytest.raises(QueueFullError):
        await backend.submit(spec)

    attempts = 0
    second_id = ""

    async def flaky(job: JobSpec) -> dict:
        nonlocal attempts, second_id
        attempts += 1
        if attempts == 1:
            # while job 1 is running (and so off the queue) a second job is admitted: the queue is
            # full again at the moment job 1 fails and needs to be re-queued
            second_id = (await backend.submit(spec)).id
            raise UpstreamError("transient")
        return {"ok": True}

    stop = asyncio.Event()
    worker = asyncio.create_task(backend.run_worker(flaky, concurrency=1, stop=stop))
    done = await backend.wait(first.id, 10)
    other = await backend.wait(second_id, 10)
    stop.set()
    await asyncio.wait_for(worker, 5)
    assert done.status == "succeeded" and done.attempts == 2, (
        "the retry was accepted although the queue was full"
    )
    assert other.status == "succeeded", "and the worker survived to run the next job"


async def test_one_unexpected_error_does_not_stop_the_worker(monkeypatch):
    backend = InProcessJobBackend(Settings(_env_file=None))
    spec = JobSpec(collection="c", path="/x")
    broken, healthy = await backend.submit(spec), await backend.submit(spec)
    original = backend._process  # noqa: SLF001
    calls = []

    async def process(job_id, handler):
        calls.append(job_id)
        if job_id == broken.id:
            raise RuntimeError("bug in bookkeeping")
        await original(job_id, handler)

    monkeypatch.setattr(backend, "_process", process)

    async def handler(job: JobSpec) -> dict:
        return {}

    stop = asyncio.Event()
    worker = asyncio.create_task(backend.run_worker(handler, concurrency=1, stop=stop))
    assert (await backend.wait(healthy.id, 10)).status == "succeeded", "the next job still runs"
    stop.set()
    await asyncio.wait_for(worker, 5)
    assert calls == [broken.id, healthy.id]


# --- embeddings -----------------------------------------------------------------------------
class SpyEmbeddings:
    def __init__(self):
        self.calls = []

    async def aembed_documents(self, texts):
        self.calls.append(("documents", list(texts)))
        return [[1.0, 0.0] for _ in texts]

    async def aembed_query(self, text):
        self.calls.append(("query", text))
        return [0.0, 1.0]


def adapter(raw, *, batch_queries):
    return LangChainEmbedder(
        "e", raw, dimensions=2, batch_size=8, service="t", bulkhead=Bulkhead(4), batch_queries=batch_queries
    )


async def test_providers_with_distinct_query_encodings_get_their_own_query_call():
    """Regression: queries were embedded with aembed_documents, so providers that encode queries
    differently (task types, instruction prefixes) silently lost recall."""
    raw = SpyEmbeddings()
    vectors = await adapter(raw, batch_queries=False).embed_queries(["q1", "q2"])
    assert sorted(c for c in raw.calls if c[0] == "query") == [("query", "q1"), ("query", "q2")]
    assert vectors == [[0.0, 1.0], [0.0, 1.0]] and not any(c[0] == "documents" for c in raw.calls)
    raw.calls.clear()
    await adapter(raw, batch_queries=False).embed_documents(["d1", "d2"])
    assert raw.calls == [("documents", ["d1", "d2"])]


async def test_symmetric_models_still_batch_queries_into_one_request():
    raw = SpyEmbeddings()
    await adapter(raw, batch_queries=True).embed_queries(["q1", "q2", "q3"])
    assert raw.calls == [("documents", ["q1", "q2", "q3"])]


async def test_shared_embedding_cache_is_keyed_by_the_model_not_the_profile_name():
    """Regression: the key used the profile name, so re-pointing a profile at a different model
    reused the old model's cached query vectors."""
    shared = MemoryCache(100)

    class Embedder:
        dimensions = 2
        model_id = "small"  # the profile name, identical in both deployments

        def __init__(self, vector):
            self.vector = vector

        async def embed_documents(self, texts):
            return [self.vector for _ in texts]

        embed_queries = embed_documents

    old = CachingEmbedder(Embedder([1.0, 0.0]), max_entries=10, shared=shared, identity="openai/model-a/")
    new = CachingEmbedder(Embedder([0.0, 1.0]), max_entries=10, shared=shared, identity="openai/model-b/")
    assert (await old.embed_queries(["same text"]))[0] == pytest.approx([1.0, 0.0])
    assert (await new.embed_queries(["same text"]))[0] == pytest.approx([0.0, 1.0]), (
        "must not see model-a's vector"
    )


# --- upload body cap --------------------------------------------------------------------------
def upload_app(limit: int = 1_000_000) -> tuple[FastAPI, list[int]]:
    app = FastAPI()
    app.add_middleware(MaxBodySizeMiddleware, path="/up", max_bytes=limit)
    reached: list[int] = []

    @app.post("/up")
    async def up(file: UploadFile = File(...)):
        reached.append(1)
        return {"bytes": len(await file.read())}

    @app.post("/other")
    async def other(payload: dict):
        return payload

    return app, reached


async def chunked_multipart(megabytes: int):
    yield b'--B\r\nContent-Disposition: form-data; name="file"; filename="a.txt"\r\nContent-Type: text/plain\r\n\r\n'
    for _ in range(megabytes * 10):
        yield b"x" * 100_000
    yield b"\r\n--B--\r\n"


async def test_oversized_uploads_are_rejected_with_413_before_reaching_the_endpoint():
    app, reached = upload_app()
    async with httpx.AsyncClient(
        transport=httpx.ASGITransport(app=app, raise_app_exceptions=False), base_url="http://t"
    ) as client:
        declared = await client.post("/up", files={"file": ("a.txt", b"x" * 2_000_000)})
        chunked = await client.post(
            "/up", content=chunked_multipart(4), headers={"Content-Type": "multipart/form-data; boundary=B"}
        )
        fine = await client.post("/up", files={"file": ("a.txt", b"x" * 1000)})
        other = await client.post("/other", json={"big": "y" * 2_000_000})
    for response in (declared, chunked):
        assert response.status_code == 413 and response.json()["code"] == "payload_too_large"
        assert "1 MB" in response.json()["detail"] or "0.95" in response.json()["detail"]
    assert fine.status_code == 200 and reached == [1], "only the small upload reached the endpoint"
    assert other.status_code == 200, "the limit applies to the upload route only"


# --- chat history budget ----------------------------------------------------------------------
def chat_service(max_chars: int) -> ChatService:
    return ChatService(
        answers=None,
        models=None,
        store=MemoryConversationStore(10, 60),
        utility_model="u",  # type: ignore[arg-type]
        settings=Settings(_env_file=None, chat_history_max_chars=max_chars),
    )


def turns(n: int, size: int) -> list[ChatMessage]:
    return [
        m
        for i in range(n)
        for m in (ChatMessage("user", f"{i}" * size), ChatMessage("assistant", f"{i}" * size))
    ]


def test_chat_history_is_trimmed_to_the_budget_oldest_turns_first():
    history = turns(10, 100)  # 2000 characters
    kept = chat_service(500)._fit(history)  # noqa: SLF001
    assert sum(len(m.content) for m in kept) <= 500 and kept == history[-len(kept) :]
    assert [m.role for m in kept][0] == "user", "trimming drops whole turns"


def test_the_latest_turn_is_always_kept_even_if_it_alone_exceeds_the_budget():
    history = turns(3, 1000)
    assert chat_service(10)._fit(history) == history[-2:]  # noqa: SLF001
    assert chat_service(10_000)._fit(history) == history  # noqa: SLF001


# --- redis timeouts, deadline helper ---------------------------------------------------------
async def test_redis_clients_carry_socket_timeouts():
    client = new_client("redis://localhost:6390/0", decode_responses=True, socket_timeout=3.0)
    kwargs = client.connection_pool.connection_kwargs
    assert kwargs["socket_timeout"] == 3.0 and kwargs["socket_connect_timeout"] == 3.0
    await client.aclose()


async def test_a_stalled_redis_fails_fast_instead_of_hanging():
    from src.runtime.cache import RedisCache

    async def accept_and_say_nothing(reader, writer):
        await asyncio.sleep(30)

    server = await asyncio.start_server(accept_and_say_nothing, "127.0.0.1", 0)
    port = server.sockets[0].getsockname()[1]
    cache = RedisCache(f"redis://127.0.0.1:{port}/0", socket_timeout=1.0)
    started = asyncio.get_running_loop().time()
    with pytest.raises(UpstreamError):
        await cache.get("k")
    assert asyncio.get_running_loop().time() - started < 5
    await cache.close()
    server.close()


def test_the_redis_socket_timeout_must_exceed_the_worker_poll_interval():
    with pytest.raises(ValidationError):
        Settings(_env_file=None, redis_socket_timeout_seconds=0.5)


async def test_deadline_turns_a_timeout_into_a_typed_error():
    with pytest.raises(RequestTimeoutError, match="0.05s"):
        async with deadline(0.05):
            await asyncio.sleep(1)
    async with deadline(1):
        await asyncio.sleep(0)
