from __future__ import annotations

import io
from pathlib import Path

import cv2
import numpy as np
import pytest
from src.chunking.ids import chunk_id
from src.chunking.keywords import KeywordExtractor
from src.chunking.recursive import RecursiveChunker
from src.core.bootstrap import build_registries
from src.core.config import Settings
from src.core.errors import ConfigError, ParseError, UnsupportedTypeError
from src.core.registry import Registries
from src.core.specs import ChunkerSpec
from src.parsing.ocr.image_ops import (
    MAX_DESKEW_DEGREES,
    decode_image,
    deskew_image,
    estimate_skew,
    limit_side,
)
from src.parsing.registry import ParserSet

from tests.helpers import make_pdf, text_png


@pytest.fixture(scope="module")
def parsers() -> ParserSet:
    settings = Settings(_env_file=None, ocr_min_image_side=64)
    return ParserSet(build_registries(settings), settings)


def units(parsers: ParserSet, path: Path):
    return list(parsers.for_path(path).iter_units(path))


# --- chunking -------------------------------------------------------------------------------
def test_chunk_ids_are_deterministic_and_content_addressed():
    base = chunk_id("/a.pdf", "text", 1, 0, "hello")
    assert base == chunk_id("/a.pdf", "text", 1, 0, "hello") and len(base) == 40
    variants = [
        chunk_id("/b.pdf", "text", 1, 0, "hello"),
        chunk_id("/a.pdf", "table", 1, 0, "hello"),
        chunk_id("/a.pdf", "text", 2, 0, "hello"),
        chunk_id("/a.pdf", "text", 1, 1, "hello"),
        chunk_id("/a.pdf", "text", 1, 0, "hello!"),
    ]
    assert len({base, *variants}) == 6, "any change of source, kind, page, index or content changes the id"


def test_recursive_chunker_splits_and_applies_the_minimum():
    chunker = RecursiveChunker(ChunkerSpec(chunk_size=100, chunk_overlap=10, min_chunk_size=30))
    assert chunker.split("") == [] and chunker.split("too short") == []
    drafts = chunker.split("Sentence number one is here. " * 20)
    assert len(drafts) > 1 and all(len(d.content) <= 100 for d in drafts)
    assert [d.index for d in drafts] == list(range(len(drafts))) and {d.total for d in drafts} == {
        len(drafts)
    }


def test_chunks_end_at_sentence_punctuation_and_never_start_with_it():
    """Regression: with the splitter's default the separator moved to the *start* of the next
    chunk, giving chunks such as '। यह चौथा वाक्य है।' and '. Third sentence...'."""
    chunker = RecursiveChunker(ChunkerSpec(chunk_size=60, chunk_overlap=0, min_chunk_size=10))
    hindi = "यह पहला वाक्य है। यह दूसरा वाक्य है। यह तीसरा वाक्य है। यह चौथा वाक्य है।"
    chunks = [d.content for d in chunker.split(hindi)]
    assert chunks == ["यह पहला वाक्य है। यह दूसरा वाक्य है। यह तीसरा वाक्य है।", "यह चौथा वाक्य है।"]
    english = [
        d.content
        for d in chunker.split(
            "First sentence is here. Second sentence is here. Third sentence is here. Fourth one."
        )
    ]
    assert english == [
        "First sentence is here. Second sentence is here.",
        "Third sentence is here. Fourth one.",
    ]


def test_keywords_work_for_a_single_text_and_for_devanagari():
    """Regression: max_df=0.95 made scikit-learn raise for one text, and the old broad except
    turned that into 'no keywords' for every single-chunk page."""
    extractor = KeywordExtractor(top_k=5)
    [single] = extractor.extract(
        ["Quarterly revenue growth exceeded expectations across cloud subscriptions"]
    )
    assert any("revenue" in k for k in single) and 0 < len(single) <= 5
    [hindi] = extractor.extract(["भारत की अर्थव्यवस्था तेजी से बढ़ रही है और निवेश बढ़ा है"])
    assert any("अर्थव्यवस्था" in k for k in hindi), "Devanagari words must not be shredded at vowel signs"


def test_keywords_edge_cases():
    extractor = KeywordExtractor(top_k=3)
    assert extractor.extract([]) == []
    assert extractor.extract(["the and of to in", "is a an the"]) == [[], []], (
        "stop-word-only batches have no vocabulary"
    )
    many = extractor.extract(["alpha beta gamma delta epsilon zeta", "alpha omega sigma"])
    assert all(len(k) <= 3 for k in many) and any("omega" in k for k in many[1])


# --- parsers --------------------------------------------------------------------------------
def test_pdf_text_tables_and_image_dedupe(parsers, tmp_path):
    png = text_png(["Scanned logo text"])
    pdf = make_pdf(
        tmp_path / "a.pdf",
        ["Revenue grew twelve percent year over year.", "Second page text."],
        table=[["Region", "Q1"], ["North", "10"]],
        image_on_page=1,
        image_png=png,
    )
    first, second = units(parsers, pdf)
    assert first.unit == 1 and "Revenue" in first.texts[0].text and first.texts[0].language == "en"
    assert first.tables and "| Region" in first.tables[0].markdown and first.tables[0].row_count == 1
    assert len(first.images) == 1 and second.images == []


def test_pdf_images_are_deduplicated_by_content_across_pages(parsers, tmp_path):
    import pymupdf

    png = text_png(["Same image everywhere"])
    pdf = make_pdf(
        tmp_path / "logo.pdf", ["page one", "page two", "page three"], image_on_page=1, image_png=png
    )
    doc = pymupdf.open(pdf)
    for page in (doc[1], doc[2]):  # separate insertions => different object numbers, identical bytes
        page.insert_image(pymupdf.Rect(72, 400, 372, 520), stream=png)
    doc.saveIncr()
    doc.close()
    assert sum(len(u.images) for u in units(parsers, pdf)) == 1


def test_pdf_ignores_tiny_images(parsers, tmp_path):
    tiny = cv2.imencode(".png", np.full((20, 20, 3), 128, np.uint8))[1].tobytes()
    pdf = make_pdf(tmp_path / "icons.pdf", ["body text"], image_on_page=1, image_png=tiny)
    assert units(parsers, pdf)[0].images == []


def test_pdf_errors_are_parse_errors_with_a_remedy(parsers, tmp_path, monkeypatch):
    broken = tmp_path / "broken.pdf"
    broken.write_bytes(b"%PDF-1.4 not really a pdf")
    with pytest.raises(ParseError, match="broken.pdf"):
        units(parsers, broken)

    import pymupdf

    secret = tmp_path / "locked.pdf"
    doc = pymupdf.open()
    doc.new_page().insert_text((72, 72), "x")
    doc.save(secret, encryption=pymupdf.PDF_ENCRYPT_AES_256, user_pw="pw", owner_pw="pw")
    with pytest.raises(ParseError, match="password"):
        units(parsers, secret)

    good = make_pdf(tmp_path / "good.pdf", ["some text"])
    monkeypatch.setattr(
        pymupdf.Page,
        "find_tables",
        lambda self, *a, **k: (_ for _ in ()).throw(RuntimeError("layout engine crashed")),
    )
    with pytest.raises(ParseError, match="PDF_EXTRACT_TABLES=false"):
        units(parsers, good)


def test_docx_sections_tables_and_images(parsers, tmp_path):
    import docx

    document = docx.Document()
    document.add_heading("Leave policy", 1)
    document.add_paragraph("Twenty days of leave.")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text, table.cell(1, 0).text, table.cell(1, 1).text = (
        "Type",
        "Days",
        "Annual",
        "20",
    )
    png = tmp_path / "pic.png"
    png.write_bytes(text_png(["Org chart"]))
    document.add_picture(str(png))
    document.add_heading("Conduct", 1)
    document.add_paragraph("Be respectful.")
    path = tmp_path / "policy.docx"
    document.save(path)
    first, second = units(parsers, path)
    assert "Twenty days" in first.texts[0].text and first.tables[0].row_count == 1 and len(first.images) == 1
    assert second.texts[0].text.startswith("Conduct")


def test_spreadsheet_and_csv_keep_headers_in_every_block(parsers, tmp_path):
    import openpyxl

    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Sales"
    sheet.append(["item", "qty"])
    for i in range(120):
        sheet.append([f"i{i}", i])
    workbook.save(tmp_path / "s.xlsx")
    (tmp_path / "s.csv").write_text("name,age\n" + "\n".join(f"p{i},{20 + i}" for i in range(75)))
    for name, expected_blocks in (("s.xlsx", 3), ("s.csv", 2)):
        blocks = units(parsers, tmp_path / name)
        assert len(blocks) == expected_blocks
        assert all(
            u.tables[0].markdown.startswith("| ") and "---" in u.tables[0].markdown.splitlines()[1]
            for u in blocks
        )
    assert "Sheet 'Sales'" in units(parsers, tmp_path / "s.xlsx")[0].tables[0].summary


def test_html_skips_scripts_and_styles_splits_sections_and_extracts_tables(parsers, tmp_path):
    page = tmp_path / "p.html"
    page.write_text(
        "<html><head><title>T</title><style>p{}</style></head><body><h1>Intro</h1><p>Hello world</p>"
        "<table><tr><th>a</th><th>b</th></tr><tr><td>1</td><td>2</td></tr></table>"
        "<h2>More</h2><p>Second &amp; last</p><script>alert('bad')</script></body></html>"
    )
    first, second = units(parsers, page)
    assert (
        first.texts[0].text == "Intro\nHello world"
        and first.tables[0].markdown.splitlines()[0] == "| a | b |"
    )
    assert second.texts[0].text == "More\nSecond & last"
    assert "alert" not in first.texts[0].text + second.texts[0].text and "T" != first.texts[0].text


def test_text_sections_and_strict_utf8(parsers, tmp_path):
    (tmp_path / "n.md").write_text("Intro paragraph.\n\n" + "Body paragraph. " * 800 + "\n\nClosing.")
    assert len(units(parsers, tmp_path / "n.md")) >= 2
    (tmp_path / "latin1.txt").write_bytes("café".encode("latin-1"))
    with pytest.raises(ParseError, match="not valid UTF-8"):
        units(parsers, tmp_path / "latin1.txt")


def test_images_single_frame_multiframe_and_corrupt(parsers, tmp_path):
    from PIL import Image

    (tmp_path / "one.png").write_bytes(text_png(["hello"]))
    [single] = units(parsers, tmp_path / "one.png")
    assert single.unit == 1 and len(single.images) == 1 and single.texts == []
    frames = [Image.fromarray(np.full((80, 80, 3), c, np.uint8)) for c in (10, 120, 240)]
    frames[0].save(tmp_path / "multi.tiff", save_all=True, append_images=frames[1:])
    assert [u.unit for u in units(parsers, tmp_path / "multi.tiff")] == [1, 2, 3]
    (tmp_path / "bad.png").write_bytes(b"not an image")
    with pytest.raises(ParseError, match="cannot read image"):
        units(parsers, tmp_path / "bad.png")


def test_parser_set_routing_rules(parsers, tmp_path):
    assert parsers.for_path(Path("x.PDF")).name == "pdf", "extensions are case-insensitive"
    with pytest.raises(UnsupportedTypeError, match="Supported:"):
        parsers.for_path(Path("x.exe"))
    with pytest.raises(UnsupportedTypeError, match="not accepted by this collection"):
        parsers.for_path(Path("x.docx"), allowed=["pdf", "text"])
    assert parsers.for_path(Path("x.pdf"), allowed=["pdf"]).name == "pdf"


def test_two_parsers_cannot_claim_the_same_extension():
    registries = build_registries(Settings(_env_file=None))

    class Rival:
        name = "rival"
        extensions = frozenset({".pdf"})

        def __init__(self, settings):
            pass

        def iter_units(self, path):
            return iter(())

    registries.parsers.register("rival", Rival)
    with pytest.raises(ConfigError, match=r"extension '\.pdf' is claimed by both"):
        ParserSet(registries, Settings(_env_file=None))


def test_a_plugin_parser_works_end_to_end_without_core_changes(tmp_path):
    from src.ports.parsing import ParsedUnit, TextBlock

    class LogParser:
        name, extensions = "log", frozenset({".log"})

        def __init__(self, settings):
            pass

        def iter_units(self, path):
            for number, line in enumerate(path.read_text().splitlines(), start=1):
                yield ParsedUnit(unit=number, texts=[TextBlock(line, "en")])

    registries = Registries()
    registries.parsers.register("log", LogParser)
    (tmp_path / "a.log").write_text("first line\nsecond line")
    pset = ParserSet(registries, Settings(_env_file=None))
    assert [u.texts[0].text for u in pset.for_path(Path("a.log")).iter_units(tmp_path / "a.log")] == [
        "first line",
        "second line",
    ]


# --- OCR image helpers ----------------------------------------------------------------------
@pytest.mark.parametrize("angle", [0, 2, -2, 5, -5, 9, -12])
def test_deskew_straightens_realistic_noisy_scans(angle):
    """Regression: the old mask (gray > 0) covered the whole page on any real scan, so skew
    was never detected and nothing was ever corrected."""
    image = decode_image(
        text_png(
            ["Quarterly revenue grew", "Operating margin improved", "Customer retention strong"], angle=angle
        )
    )
    assert abs(estimate_skew(image) + angle) <= 0.2, "measured skew is the negative of the applied rotation"
    assert abs(estimate_skew(deskew_image(image))) <= 0.5


def test_deskew_leaves_implausible_angles_and_blank_pages_alone():
    blank = np.full((200, 300, 3), 255, np.uint8)
    assert deskew_image(blank) is blank
    assert abs(MAX_DESKEW_DEGREES) == 15.0


def test_decode_and_downscale():
    with pytest.raises(ParseError, match="could not be decoded"):
        decode_image(b"garbage")
    big = np.zeros((4000, 1000, 3), np.uint8)
    small = limit_side(big, 2000)
    assert max(small.shape[:2]) == 2000 and limit_side(small, 2000) is small


def test_image_io_roundtrip_helper():
    buffer = io.BytesIO(text_png(["x"]))
    assert buffer.getbuffer().nbytes > 100
